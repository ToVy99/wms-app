"""Qt offscreen smoke test for the IntraCity 3.4 native Check Prod UI.

Tests require PySide6 and deliberately avoid WMS login or network requests.
Run: set QT_QPA_PLATFORM=offscreen && python -m unittest test_prod_native_layout
"""
import os
import copy
import tempfile
import unittest
from unittest.mock import patch
from PySide6.QtCore import QObject, Signal, Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QTabBar, QLabel, QPushButton, QLineEdit
from productivity_tab import CheckProdTab, TABS


class FakeBridge(QObject):
    completed = Signal(str, str)
    sessionChanged = Signal(bool, str)
    connected = False
    session_message = 'Offline UI test'

    def request(self, *_args):
        raise AssertionError('UI layout smoke test must not call WMS')


class NativeProdLayoutTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        self.bridge = FakeBridge()
        self.widget = CheckProdTab(self.bridge, {}, lambda _config: True)
        self.widget.show()
        self.app.processEvents()

    def tearDown(self):
        self.widget.close()
        self.widget.deleteLater()
        self.app.processEvents()

    def test_all_work_tabs_remain_exposed(self):
        tabbar = self.widget.tab_select
        self.assertEqual([tabbar.tabText(i) for i in range(tabbar.count())], list(TABS))
        self.assertFalse(tabbar.usesScrollButtons())
        width = tabbar.sizeHint().width()
        for index in range(tabbar.count()):
            tabbar.setCurrentIndex(index)
            self.app.processEvents()
            self.assertEqual(tabbar.sizeHint().width(), width)

    def test_no_duplicate_connection_strip_above_controls(self):
        all_labels=[label.text() for label in self.widget.findChildren(QLabel)]
        self.assertFalse(any('Dữ liệu chung từ Chrome WMS' in text for text in all_labels))
        self.assertFalse(any('Chưa kết nối WMS' in text for text in all_labels))
        self.assertEqual(self.widget._wms_connected,False)
        self.bridge.sessionChanged.emit(True,'Đang dùng chung tab WMS')
        self.assertTrue(self.widget._wms_connected)
        self.assertEqual(self.widget._wms_connection_message,'Đang dùng chung tab WMS')

    def test_target_has_no_long_unit_label_and_all_modes_fixed(self):
        texts = [label.text() for label in self.widget.findChildren(QLabel)]
        self.assertNotIn('Items/h hoặc Orders/h', texts)
        buttons = (self.widget.bau,self.widget.cp,self.widget.item,self.widget.order,
                   self.widget.shift1,self.widget.shift2)
        for btn in buttons:
            self.assertEqual(btn.width(),68)
            self.assertEqual(btn.height(),28)

    def test_dark_light_styling_and_metric_icons(self):
        icons=[label for label in self.widget.findChildren(QLabel) if label.objectName()=='prodMetricIcon']
        self.assertEqual(len(icons), 5)
        for light in (True,False,True):
            self.widget.set_theme(light)
            self.assertEqual(self.widget._light_theme,light)
            self.assertIn('QFrame#prodMetricCard',self.widget.styleSheet())
            self.assertIn('QPushButton[prodChoice="true"]:checked', self.widget.styleSheet())
        self.assertEqual(set(self.widget.metric_values),{'qty','weighted','avg','manHour','below'})
        self.assertFalse(self.bridge.connected)

    def load_staff(self,tab):
        w=self.widget
        w.tab_select.setCurrentIndex(TABS.index(tab))
        base=500 if tab=='Check&Pack' else 1
        w.rows=[dict(name=name,email=email,operatorName=name,operatorEmail=email,
                     stationId=f'CS{base+i:05d}',stationName=f'Station {i}',
                     productivity=prod,prod=prod,items=prod,orders=prod/2,
                     manHour=1,staffProd=prod,hPrev=prod,hCur=prod)
                for i,(name,email,prod) in enumerate([
                    ('An','an@example.com',10),('An',' AN@example.com ',20),
                    ('An','another@example.com',30),('Bình','binh@example.com',40)])]
        w.cache_key=w._view_key();w.show_all.setChecked(True);w._render()
        w.resize(1250,850);self.app.processEvents()

    def click_cell(self,row,col):
        table=self.widget.table;item=table.item(row,col)
        table.scrollToItem(item);self.app.processEvents()
        QTest.mouseClick(table.viewport(),Qt.LeftButton,pos=table.visualItemRect(item).center())
        self.app.processEvents()

    def row_for_email(self,email):
        w=self.widget;col=1 if w.tab in ('Sort','Check','Pack') else 6
        return next(r for r in range(w.table.rowCount()) if w.table.item(r,col).text()==email)

    def assert_highlighted(self,email):
        w=self.widget;email_col=1 if w.tab in ('Sort','Check','Pack') else 6
        for row in range(w.table.rowCount()):
            expected=email is not None and w.table.item(row,email_col).text().strip().casefold()==email
            for col in range(w.table.columnCount()):
                item=w.table.item(row,col)
                self.assertEqual(item.font().bold(),expected)
                self.assertEqual(item.data(Qt.ItemDataRole.BackgroundRole) is not None,expected)

    def test_name_click_groups_all_rows_toggles_and_switches_employee_in_every_tab(self):
        for tab in TABS:
            with self.subTest(tab=tab):
                self.load_staff(tab)
                w=self.widget;name_col=0 if tab in ('Sort','Check','Pack') else 5
                before=copy.deepcopy(w.rows)
                metrics={key:value.text() for key,value in w.metric_values.items()}
                self.click_cell(self.row_for_email('an@example.com'),name_col)
                self.assert_highlighted('an@example.com')
                self.click_cell(self.row_for_email(' AN@example.com '),name_col)
                self.assert_highlighted(None)
                self.click_cell(self.row_for_email('an@example.com'),name_col)
                self.click_cell(self.row_for_email('binh@example.com'),name_col)
                self.assert_highlighted('binh@example.com')
                QTest.mouseClick(w.clear_highlight,Qt.LeftButton)
                self.assert_highlighted(None)
                self.assertFalse(w.clear_highlight.isEnabled())
                self.assertEqual(w.rows,before)
                self.assertEqual({key:value.text() for key,value in w.metric_values.items()},metrics)

    def test_highlight_survives_sort_filter_refresh_theme_and_preserves_export(self):
        from openpyxl import load_workbook
        self.load_staff('Check Station');w=self.widget
        self.click_cell(self.row_for_email('an@example.com'),5)
        w._flip();w.show_all.setChecked(False);w._render()
        for light in (True,False):
            w.set_theme(light);self.assert_highlighted('an@example.com')
        expected=[[w.table.item(r,c).text() for c in range(w.table.columnCount())]
                  for r in range(w.table.rowCount())]
        with tempfile.TemporaryDirectory() as td:
            path=td+'/prod.xlsx'
            with patch('productivity_tab.QFileDialog.getSaveFileName',return_value=(path,'')):
                w._export()
            book=load_workbook(path)
            self.assertEqual([list(row) for row in book.active.iter_rows(min_row=2,values_only=True)],expected)
            book.close()
        w.tab_select.setCurrentIndex(TABS.index('Pack'))
        self.assertEqual(w.table.rowCount(),0)
        self.assertFalse(w.clear_highlight.isEnabled())

    def test_unknown_employee_and_non_name_cells_do_not_highlight(self):
        self.load_staff('Check Station');w=self.widget
        self.click_cell(0,0);self.assert_highlighted(None)
        w.rows=[dict(stationId='CS00001',prod=10,manHour=1,items=10)]
        w._render();self.click_cell(0,5);self.assert_highlighted(None)

    def test_email_input_removed_but_existing_api_email_preserved(self):
        w=self.widget
        self.assertFalse(any('Email WMS' in field.placeholderText() for field in w.findChildren(QLineEdit)))
        self.assertFalse(hasattr(w,'email'))
        for email in ('',' saved@example.com '):
            with self.subTest(email=email):
                w.prefs['loginEmail']=email;self.bridge.connected=True
                with patch.object(w,'_load',return_value=iter([])) as load:
                    w._start()
                self.assertEqual(load.call_args.args[3],email.strip())
                self.assertEqual(w.prefs['loginEmail'],email)


if __name__ == '__main__':
    unittest.main()
