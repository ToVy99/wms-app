"""Native implementation of the IntraCity web 1.4 data rules."""
import copy
import io
import json
import unicodedata
import zipfile
from datetime import datetime, timedelta
from openpyxl import load_workbook

ORIGIN = 'https://wms.ssc.shopee.vn'
URL_ORDER = ORIGIN + '/api/v2/apps/process/outbound/salesorder/search_order'
URL_EXPORT = ORIGIN + '/api/v2/apps/basic/reportcenter/create_export_task'
URL_EXPORT_LIST = ORIGIN + '/api/v2/apps/basic/reportcenter/search_export_task?pageno=1&count=20&export_module=2&task_type=722&is_myself=1'
STATUSES = ['Created','Pending Pick','Picking','Picked','Pick Fail','Checking','Checked','Pre Sorting','Pre Sorted','Sorting','Sorted','Packing','Packed','Shipping','Outbound','Cancel']
DEFAULT_CONFIG = {
 'intra': {'name':'Intra City','channels':['50057','50021'],'cots':[{'id':str(i+1),'name':f'COT {i+1}','start':a,'end':b} for i,(a,b) in enumerate([('20:00','23:00'),('23:00','02:00'),('02:00','16:00'),('16:00','20:00')])]},
 'aha': {'name':'AhaMove','channels':['50033','50044'],'cots':[{'id':str(i+1),'name':f'COT {i+1}','start':a,'end':b} for i,(a,b) in enumerate([('08:00','13:00'),('13:00','18:00'),('18:00','08:00')])]},
 'sdd': {'name':'SDD','channels':['50051'],'cots':[{'id':str(i+1),'name':f'COT {i+1}','start':a,'end':b} for i,(a,b) in enumerate([('18:00','04:00'),('04:00','09:00'),('09:00','13:30'),('13:30','18:00')])]},
 'spxck': {'name':'SPX CK (Cồng kềnh)','channels':['50025'],'previous':True,'cots':[{'id':'daily','name':'17H → 17H','start':'17:00','end':'17:00'}]},
}

def time_range(day, carrier, cot):
    beg = datetime.strptime(day+' '+cot['start'], '%Y-%m-%d %H:%M')
    end = datetime.strptime(day+' '+cot['end'], '%Y-%m-%d %H:%M')
    if carrier.get('previous'):
        beg -= timedelta(days=1)
    elif cot['end'] <= cot['start']:
        end += timedelta(days=1)
    return beg, end

def export_body(beg, end, channels):
    if not channels:
        raise ValueError("Chưa có mã nhà vận chuyển cho nhóm COT.")
    a,b = int(beg.timestamp()),int(end.timestamp())
    extra = dict(beg_ctime=a,end_ctime=b,channel_id_list=list(channels),from_listpage=1,include_sku_list=1,order_type=0,date_ref=0,time_from=a,time_to=b)
    return dict(export_module=2,task_type=722,extra_data=json.dumps(extra,separators=(',',':')))

def choose_job(jobs, old_ids, started, task_id=None):
    if task_id is not None:
        return next((j for j in jobs if str(j.get('task_id'))==str(task_id)),None)
    fresh = [j for j in jobs if str(j.get('task_id')) not in old_ids and float(j.get('ctime') or 0)>=started-10]
    if len(fresh)>1:
        raise ValueError('Có nhiều Export cùng lúc; đã dừng để tránh lấy nhầm report.')
    return fresh[0] if fresh else None

def norm(v):
    return unicodedata.normalize('NFC',str(v or '').strip().lower())

def carrier_match(key, service, state):
    v = norm(service)
    if key=='intra':
        return v in ('spx express','spx express ndd - trong ngày') and norm(state) in ('hà nội','thành phố hà nội')
    if key=='aha':
        return v in ('ahamove','ahamove sbs') or v.startswith(('ahamove -','ahamove sbs -'))
    if key=='sdd':
        return v=='spx express sbs trong ngày' or v.startswith(('spx express sbs trong ngày -','spx express sbs - trong ngày'))
    if key=='spxck':
        return v=='spx - hàng cồng kềnh'
    raise ValueError('Không có bộ lọc nhóm '+key)

def report_files(data, name):
    """Yield each workbook once, including workbooks in ZIP subfolders."""
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if 'xl/workbook.xml' in archive.namelist():
            yield name, data
            return
        members=[m for m in archive.infolist() if not m.is_dir()
                 and m.filename.lower().endswith('.xlsx')
                 and not m.filename.startswith('__MACOSX/')
                 and not m.filename.rsplit('/',1)[-1].startswith(('~$','._'))]
        if not members:
            raise ValueError('ZIP không có file XLSX.')
        for member in members:
            yield member.filename, archive.read(member)


def report_sheets(data):
    """Read sheets one at a time; retain OpenPyXL fallback for malformed exports."""
    book=None; fallback=None
    try:
        try:
            import pandas as pd
            book=pd.ExcelFile(io.BytesIO(data),engine='calamine')
            names=book.sheet_names
        except Exception:
            fallback=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
            names=fallback.sheetnames
        for name in names:
            if book is not None:
                try:
                    frame=book.parse(name,header=None,dtype=object,keep_default_na=False)
                    yield name, frame.itertuples(index=False,name=None)
                    del frame
                    continue
                except Exception:
                    if fallback is None:
                        fallback=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
            sheet=fallback[name];sheet.reset_dimensions()
            yield name, sheet.iter_rows(values_only=True)
    finally:
        if book is not None:book.close()
        if fallback is not None:fallback.close()


def parse_report(data, name, key):
    grouped={};raw=filtered=files=sheets=0
    required=['WMS Order No','Status','New 3PL']+(['Buyer State'] if key=='intra' else [])
    for filename, workbook in report_files(data,name):
        files+=1;valid_sheets=0
        try:
            for sheetname, rows in report_sheets(workbook):
                headers=next(rows,None)
                if headers is None:continue
                indices={str(v or '').strip():i for i,v in enumerate(headers)}
                # Ignore notes/summary sheets, but never silently skip a broken order sheet.
                if 'WMS Order No' not in indices:continue
                for col in required:
                    if col not in indices:raise ValueError(f'Sheet {sheetname}: Report thiếu {col}')
                valid_sheets+=1;sheets+=1
                for row in rows:
                    def get(col):
                        i = indices.get(col)
                        return row[i] if i is not None and i<len(row) and row[i] is not None else ''
                    ob = str(get('WMS Order No')).strip()
                    if not ob: continue
                    raw += 1
                    if not carrier_match(key,get('New 3PL'),get('Buyer State')): continue
                    filtered += 1
                    if ob not in grouped:
                        status = str(get('Status')).strip()
                        code = next((i for i,s in enumerate(STATUSES) if s.lower()==status.lower()),status)
                        grouped[ob] = dict(order_number=ob,order_status=code,status_text=status,wave_type=str(get('Wave Type')).strip(),sns=set(),picking=set(),bsks=set(),sku_names=set(),search_values=set(),area='-',ctime_text=str(get('Create Time')),cutoff_text=str(get('Cut off Time')))
                    o = grouped[ob]
                    # Keep all report columns searchable, including fields not shown in the table.
                    o['search_values'].update(str(v).strip() for v in row if v is not None and str(v).strip())
                    for col,dest in [('Shopee order SN','sns'),('Picking ID','picking'),('Device ID','bsks'),('Basket ID','bsks'),('SKU Name','sku_names')]:
                        v = get(col)
                        if v: o[dest].add(str(v))
            if not valid_sheets:
                raise ValueError('Report thiếu WMS Order No hoặc không có sheet dữ liệu đơn.')
        except Exception as exc:
            raise ValueError(f'{filename}: {exc}') from exc
    return dict(orders=[o for o in grouped.values() if o['order_status']!=15],
                rawRows=raw,filteredRows=filtered,reportFiles=files,reportSheets=sheets)

def status_name(code):
    return STATUSES[code] if isinstance(code,int) and 0<=code<len(STATUSES) else str(code)

def progress(orders):
    done = sum(o['order_status']==14 for o in orders)
    return done,len(orders),100*done/len(orders) if orders else 0

