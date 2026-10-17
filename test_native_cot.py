"""Regression tests for native COT, full export, and Chrome bridge."""
import base64
import copy
import io
import json
import os
import time
import unittest
import zipfile
from pathlib import Path
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
from openpyxl import Workbook
from PySide6.QtWidgets import QApplication
from PySide6.QtCore import QDate
from cot_logic import *
from chrome_cot import ChromeCotTab, NativeBridge, validate_request, FETCH_JS


def fixture(rows):
    wb=Workbook(); ws=wb.active
    ws.append(['WMS Order No','Status','New 3PL','Buyer State','Shopee order SN','Wave Type','Picking ID','Device ID','Basket ID','Create Time','Cut off Time'])
    for row in rows: ws.append(row)
    out=io.BytesIO();wb.save(out);return out.getvalue()

class NativeTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls): cls.app=QApplication.instance() or QApplication([])
    def setUp(self): self.tab=ChromeCotTab({},lambda c:True)
    def tearDown(self):
        if self.tab.parser: self.tab.parser.wait(10000)
        import shiboken6
        shiboken6.delete(self.tab);self.app.processEvents()
    def wait(self,predicate):
        end=time.monotonic()+10
        while not predicate() and time.monotonic()<end:
            self.app.processEvents();time.sleep(.005)
        self.assertTrue(predicate())
    def test_startup_and_scope_changes_never_call_wms(self):
        calls=[];self.tab.request=lambda m,c:calls.append(m)
        self.tab.carrier.setCurrentIndex(1);self.tab.date.setDate(QDate(2026,12,31));self.tab.select_cot('2')
        self.assertEqual(calls,[]);self.assertIsNone(self.tab.bridge.worker);self.assertEqual(self.tab.carrier.count(),4)

    def test_grouped_area_counts_filter_and_missing_are_local(self):
        self.tab.orders=parse_report(fixture([[f'OB{i}','Picked','SPX Express','Hà Nội',f'SN{i}','Single',f'P{i}'] for i in range(7)]),'x.xlsx','intra')['orders']
        for o,area in zip(self.tab.orders,['A1','A2','D1','LC1','CF','A1,A2,LC2','-']):o['area']=area
        self.tab.active_status=3;self.tab.render_all()
        self.assertIn('A: 3',self.tab.area_summary.text());self.assertIn('LC: 2',self.tab.area_summary.text());self.assertIn('CF: 1',self.tab.area_summary.text())
        self.assertIn('Chưa có Area: 1',self.tab.area_summary.text())
        calls=[];self.tab.request=lambda m,c:calls.append(m)
        self.tab.area.setCurrentIndex(self.tab.area.findData('A'));self.assertEqual(self.tab.table.rowCount(),3)
        self.tab.area.setCurrentIndex(self.tab.area.findData('LC'));self.assertEqual(self.tab.table.rowCount(),2)
        self.tab.area.setCurrentIndex(self.tab.area.findData('-'));self.assertEqual(self.tab.table.rowCount(),1)
        self.assertEqual(self.tab.total_card.value.text(),'7');self.assertEqual(calls,[])

    def test_sku_names_aggregate_without_duplicating_orders(self):
        wb=Workbook();ws=wb.active;ws.append(['WMS Order No','Status','New 3PL','Buyer State','SKU Name'])
        ws.append(['OB1','Picked','SPX Express','Hà Nội','Áo xanh']);ws.append(['OB1','Picked','SPX Express','Hà Nội','Quần đen']);ws.append(['OB1','Picked','SPX Express','Hà Nội','Áo xanh'])
        out=io.BytesIO();wb.save(out)
        self.tab.orders=parse_report(out.getvalue(),'x.xlsx','intra')['orders'];self.tab.render_all()
        self.assertEqual(len(self.tab.orders),1);self.assertEqual(self.tab.orders[0]['sku_names'],{'Áo xanh','Quần đen'})
        self.assertIn('Áo xanh',self.tab.table.item(0,6).text());self.assertIn('Quần đen',self.tab.table.item(0,6).text())
        self.tab.search.setText('quan den');self.assertEqual(self.tab.table.rowCount(),1)

    def test_search_all_report_fields_is_local_and_keeps_totals(self):
        raw=fixture([
            ['OB1','Picked','SPX Express','Hà Nội','SN1','Single','PICK-A','DEVICE-1','BASKET-1'],
            ['OB1','Picked','SPX Express','Hà Nội','SN-EXTRA','Single','PICK-EXTRA','DEVICE-EXTRA','BASKET-EXTRA'],
            ['OB2','Outbound','SPX Express','Hà Nội','SN2','Multi','PICK-B','DEVICE-2','BASKET-2'],
        ])
        self.tab.orders=parse_report(raw,'search.xlsx','intra')['orders']
        self.tab.orders[0]['area']='Khu Đóng Gói'
        self.tab.render_all()
        calls=[];self.tab.request=lambda m,c:calls.append(m)
        for term in ['ob1','sn-extra','pick-extra','basket-extra','device-extra','ha noi','khu dong goi','picked pick-a']:
            self.tab.search.setText(term)
            self.assertIn('OB1',[o['order_number'] for o in self.tab.visible()])
        self.tab.select_visible();self.assertTrue(self.tab.selected)
        self.tab.search.setText('pick-b');self.assertFalse(self.tab.selected)
        self.assertEqual([o['order_number'] for o in self.tab.visible()],['OB2'])
        self.tab.copy_obvn();self.assertEqual(self.app.clipboard().text(),'OB2')
        self.assertEqual(self.tab.total_card.value.text(),'2')
        self.assertEqual(self.tab.percent_card.value.text(),'50.0%')
        self.tab.search.setText('khong-ton-tai');self.assertEqual(self.tab.table.rowCount(),0)
        self.tab.search.clear();self.assertEqual(self.tab.table.rowCount(),2)
        self.tab.toggle_picked_baskets(True);self.tab.search.setText('pick-b');self.assertEqual(self.tab.table.rowCount(),0)
        self.tab.search.setText('pick-a');self.assertEqual(self.tab.table.rowCount(),4)
        self.tab.select_cot('2');self.assertEqual(self.tab.search.text(),'pick-a')
        self.assertEqual(calls,[])
    def test_exact_windows_and_year_rollover(self):
        expected={'intra':[('20:00','23:00'),('23:00','02:00'),('02:00','16:00'),('16:00','20:00')],'aha':[('08:00','13:00'),('13:00','18:00'),('18:00','08:00')],'sdd':[('18:00','04:00'),('04:00','09:00'),('09:00','13:30'),('13:30','18:00')],'spxck':[('17:00','17:00')]}
        for k,c in DEFAULT_CONFIG.items(): self.assertEqual([(x['start'],x['end']) for x in c['cots']],expected[k])
        c=DEFAULT_CONFIG['intra'];a,b=time_range('2026-12-31',c,c['cots'][1]);self.assertEqual((b-a).total_seconds(),10800);self.assertEqual(b.year,2027);self.assertEqual(b.minute,0)
        c=DEFAULT_CONFIG['spxck'];a,b=time_range('2026-01-01',c,c['cots'][0]);self.assertEqual((b-a).total_seconds(),86400);self.assertEqual(a.year,2025)
    def test_picked_seven_dedupe_cancel_wave_and_bsk(self):
        rows=[[f'OB{i}','Picked','SPX Express','Hà Nội',f'SN{i}','Single' if i<4 else 'Multi',f'P{i}',f'D{i}',f'B{i}'] for i in range(7)]
        rows+=[['OB7','Outbound','SPX Express NDD - Trong Ngày','Thành phố Hà Nội'],['C','Cancel','SPX Express','Hà Nội'],['WRONG','Picked','SPX Express','Bắc Ninh'],['WRONG3','Picked','Ahamove','Hà Nội'],rows[0]]
        self.tab.orders=parse_report(fixture(rows),'x.xlsx','intra')['orders'];self.tab.active_status=3;self.tab.render_all()
        self.assertEqual(progress(self.tab.orders),(1,8,12.5));self.assertEqual(self.tab.table.rowCount(),7)
        self.assertIn('57.14%',self.tab.wave.itemText(1));self.assertEqual(self.tab.orders[0]['bsks'],{'D0','B0'})
        self.tab.wave.setCurrentIndex(self.tab.wave.findData('Single'));self.assertEqual(self.tab.table.rowCount(),4)
        self.tab.select_visible();self.tab.copy_obvn();self.assertEqual(len(self.app.clipboard().text().splitlines()),4)
    def test_all_carrier_filters_empty_and_outer_zip(self):
        rows=[['A','Picked','Ahamove - Trong Ngày','Bắc Ninh'],['B','Picked','Ahamove SBS - Trong Ngày','Đà Nẵng'],['S','Picked','SPX Express SBS - Trong Ngày','Bắc Ninh'],['K','Picked','SPX - Hàng Cồng Kềnh','Đà Nẵng'],['C','Cancel','Ahamove','Bắc Ninh']]
        raw=fixture(rows);out=io.BytesIO()
        with zipfile.ZipFile(out,'w') as z:z.writestr('report.xlsx',raw)
        for k,expected in [('intra',[]),('aha',['A','B']),('sdd',['S']),('spxck',['K'])]:self.assertEqual([o['order_number'] for o in parse_report(out.getvalue(),'x.zip',k)['orders']],expected)
        self.assertEqual(parse_report(fixture([]),'x.xlsx','intra')['orders'],[])
    def test_wms_incorrect_a1_dimension_reads_all_columns_and_rows(self):
        import re
        from openpyxl import load_workbook
        original=fixture([[f'OB{i}','Picked','SPX Express','Hà Nội',f'SN{i}'] for i in range(2112)])
        out=io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(original)) as src, zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as dst:
            for item in src.infolist():
                data=src.read(item.filename)
                if item.filename=='xl/worksheets/sheet1.xml':
                    data=re.sub(rb'<dimension ref="[^"]+"',b'<dimension ref="A1"',data)
                dst.writestr(item,data)
        broken=out.getvalue()
        wb=load_workbook(io.BytesIO(broken),read_only=True,data_only=True)
        self.assertEqual(next(wb.active.iter_rows(values_only=True)),('WMS Order No',))
        wb.close()
        parsed=parse_report(broken,'wms.xlsx','intra')
        self.assertEqual(len(parsed['orders']),2112)
        self.assertEqual(parsed['rawRows'],2112)
        self.assertEqual(parsed['orders'][-1]['sns'],{'SN2111'})

    def test_2112_orders_one_export_no_second_pass(self):
        data=fixture([[f'OB{i}','Picked','SPX Express','Hà Nội'] for i in range(2112)])
        calls=[]
        def request(msg,callback):
            calls.append(msg)
            if msg['type']=='WMS_DOWNLOAD':callback(base64.b64encode(data).decode())
            elif msg['type']=='WMS_POST':callback({'task_id':'NEW'})
            elif len(calls)==1:callback({'list':[]})
            else:callback({'list':[dict(task_id='NEW',ctime=int(time.time()),task_status=2,download_link=ORIGIN+'/oss_downloads_v2/fixture',export_file_name='fixture.xlsx')]})
        self.tab.bridge.connected=True;self.tab.request=request;self.tab.fetch_dataset();self.wait(lambda:not self.tab.busy)
        self.assertEqual(len(self.tab.orders),2112);self.assertEqual(len(calls),4)
        self.assertEqual(sum(m['url']==URL_EXPORT for m in calls),1);self.assertFalse(any(m['url']==URL_ORDER for m in calls))
        extra=json.loads(calls[1]['body']['extra_data']);self.assertEqual(extra['date_ref'],0);self.assertEqual(extra['include_sku_list'],1)
    def test_export_filters_all_carriers_and_keeps_saved_windows(self):
        saved=copy.deepcopy(DEFAULT_CONFIG)
        saved['intra']['channels']=[]
        saved['intra']['cots'][0]['start']='19:30'
        tab=ChromeCotTab({'cot_web_v14':{'intracityConfigV11':saved}},lambda c:True)
        try:
            self.assertEqual(tab.config['intra']['cots'][0]['start'],'19:30')
            tab.bridge.connected=True
            expected={'intra':['50057','50021'],'aha':['50033','50044'],'sdd':['50051'],'spxck':['50025']}
            for key,channels in expected.items():
                calls=[]
                def request(msg,cb):
                    calls.append(msg)
                    if msg['type']=='WMS_POST':cb({'task_id':'NEW'})
                    elif msg['type']=='WMS_DOWNLOAD':cb(base64.b64encode(fixture([])).decode())
                    elif len(calls)==1:cb({'list':[]})
                    else:cb({'list':[dict(task_id='NEW',task_status=2,download_link=ORIGIN+'/oss_downloads_v2/fixture',export_file_name='fixture.xlsx')]})
                tab.request=request
                tab.carrier.setCurrentIndex(tab.carrier.findData(key))
                self.assertEqual(calls,[])
                tab.fetch_dataset();self.wait(lambda:not tab.busy)
                posts=[m for m in calls if m['type']=='WMS_POST']
                self.assertEqual(len(posts),1)
                self.assertEqual(sum(m['type']=='WMS_DOWNLOAD' for m in calls),1)
                extra=json.loads(posts[0]['body']['extra_data'])
                self.assertEqual(extra['channel_id_list'],channels)
                _,_,beg,end=tab.current_range()
                self.assertEqual(extra['beg_ctime'],int(beg.timestamp()))
                self.assertEqual(extra['end_ctime'],int(end.timestamp()))
                self.assertEqual(extra['time_from'],extra['beg_ctime'])
                self.assertEqual(extra['time_to'],extra['end_ctime'])
        finally:
            if tab.parser:tab.parser.wait(10000)
            import shiboken6
            shiboken6.delete(tab)

    def test_status_area_200_cached_and_filters_local(self):
        self.tab.orders=parse_report(fixture([[f'OB{i}','Picked','SPX Express','Hà Nội'] for i in range(7)]),'x.xlsx','intra')['orders'];self.tab.render_all();calls=[]
        def request(msg,cb):calls.append(msg);cb({'total':7,'list':[dict(order_number=f'OB{i}',pre_hit_zone_list=['Z1']) for i in range(7)]})
        self.tab.request=request;self.tab.select_status(3)
        self.assertEqual(self.tab.table.rowCount(),7);self.assertEqual(calls[0]['body']['count'],200)
        self.tab.select_status(None);self.tab.select_status(3);self.tab.area.setCurrentIndex(1);self.assertEqual(len(calls),1)
    def test_429_stops_no_retry_and_ambiguous_jobs_rejected(self):
        with self.assertRaises(ValueError):choose_job([dict(task_id='a',ctime=100),dict(task_id='b',ctime=100)],set(),100)
        self.assertIsNone(choose_job([dict(task_id='old',ctime=100)],{'old'},100))
        self.tab.bridge.connected=True;calls=[]
        def request(rid,raw):calls.append(json.loads(raw));self.tab.bridge.completed.emit(rid,json.dumps(dict(success=False,error='HTTP 429')))
        self.tab.bridge.request=request;self.tab.fetch_dataset();self.assertFalse(self.tab.busy);self.assertEqual(len(calls),1);self.assertIn('429',self.tab.message.text());self.assertIn('resp.status===429',FETCH_JS)
    def test_picked_basket_filter_is_local_and_counts_stay_global(self):
        rows=[['A','Picked','SPX Express','Hà Nội','SNA','Single','PA','BSK-A',''],
              ['B','Picked','SPX Express','Hà Nội','SNB','Multi','PB','','BSK-B'],
              ['C','Picked','SPX Express','Hà Nội','SNC','Single','PC','',''],
              ['D','Picked','SPX Express','Hà Nội','SND','Single','PD',' - ','null'],
              ['E','Outbound','SPX Express','Hà Nội','SNE','Single','PE','BSK-E',''],
              ['A','Picked','SPX Express','Hà Nội','SNA2','Single','PA','','BSK-A2']]
        self.tab.orders=parse_report(fixture(rows),'baskets.xlsx','intra')['orders']; self.tab.render_all()
        calls=[]; self.tab.request=lambda m,c:calls.append(m)
        self.tab.selected={'E'}; self.tab.active_area='OTHER'; self.tab.active_wave='OTHER'
        self.tab.basket_filter.click()
        self.assertEqual([o['order_number'] for o in self.tab.visible()],['A','A','B'])
        self.assertEqual(self.tab.table.item(0,4).text(),'BSK-A')
        self.assertEqual(self.tab.table.rowCount(),3); self.assertEqual(self.tab.basket_filter.text(),'Picked có mã rổ')
        self.assertEqual(self.tab.selected,set()); self.tab.copy_obvn(); self.assertEqual(self.app.clipboard().text(),'A\nB')
        self.assertEqual(self.tab.total_card.value.text(),'5'); self.assertEqual(self.tab.outbound_card.value.text(),'1'); self.assertEqual(self.tab.percent_card.value.text(),'20.0%')
        self.tab.basket_filter.click(); self.assertEqual(self.tab.table.rowCount(),4)
        self.assertEqual(calls,[])
        self.tab.scope_changed(); self.assertFalse(self.tab.picked_baskets_only); self.assertEqual(self.tab.basket_filter.text(),'Picked có mã rổ')

    def test_status_panel_keeps_slots_and_height_across_data_changes(self):
        self.tab.resize(1500,800); self.tab.show(); self.app.processEvents()
        buttons=dict(self.tab.status_buttons)
        slots={code:self.tab.status_row.getItemPosition(self.tab.status_row.indexOf(button)) for code,button in buttons.items()}
        height=self.tab.status_panel.height()
        for statuses in [[3]*7+[14]*20, [0]*30+[14], []]:
            self.tab.orders=[dict(order_status=status,bsks=set(),area='-',wave_type='-',order_number=str(i),sns=set(),ctime_text='',cutoff_text='') for i,status in enumerate(statuses)]
            self.tab.render_all(); self.app.processEvents()
            self.assertEqual(self.tab.status_panel.height(),height)
            self.assertEqual(self.tab.status_row.count(),13)
            for code,button in buttons.items():
                self.assertIs(self.tab.status_buttons[code],button)
                self.assertEqual(self.tab.status_row.getItemPosition(self.tab.status_row.indexOf(button)),slots[code])
                self.assertTrue(self.tab.status_panel.rect().contains(button.geometry()))
            self.assertEqual(buttons[3].number.text(),str(statuses.count(3)))
            self.assertEqual(buttons[14].number.text(),str(statuses.count(14)))

    def test_disconnected_bridge_and_url_validation(self):
        bridge=NativeBridge({},lambda c:True);results=[];bridge.completed.connect(lambda rid,raw:results.append(json.loads(raw)))
        bridge.request('1',json.dumps(dict(type='WMS_GET',url=URL_EXPORT_LIST)));self.assertFalse(results[-1]['success']);self.assertIsNone(bridge.worker)
        for url in ['http://wms.ssc.shopee.vn/api/v2/apps/basic/reportcenter/search_export_task','https://wms.ssc.shopee.vn.evil.example/','https://wms.ssc.shopee.vn/api/unknown']:
            with self.assertRaises(ValueError):validate_request(dict(type='WMS_GET',url=url))
    def test_ui_has_no_embedded_browser(self):
        self.assertFalse(hasattr(self.tab,'view'))
        code=(Path(__file__).parent/'chrome_cot.py').read_text(encoding='utf-8')
        self.assertNotIn('QtWebEngine',code);self.assertNotIn('urllib.request',code)

if __name__=='__main__':unittest.main(verbosity=2)


