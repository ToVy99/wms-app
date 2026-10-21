"""Native PySide6 Check Prod for IntraCity 3.4 Beta.

UI is implemented with the same Qt widgets as IntraCity, not Chromium/WebEngine.
The calculation routines are translated from WMS Productivity core.js and
popup.js into prod_native_core.py, with JavaScript parity fixture tests.
All WMS requests reuse the authenticated Chrome/extension bridge unchanged.
"""
import json
import math
import re
import time
from datetime import datetime, timedelta, timezone
from pathlib import Path
from uuid import uuid4

from PySide6.QtCore import Qt, QDate, QTime, QTimer
from PySide6.QtGui import QColor, QFont
from PySide6.QtWidgets import (QWidget, QVBoxLayout, QHBoxLayout, QGridLayout, QLabel,
    QPushButton, QFrame, QComboBox, QCheckBox, QLineEdit, QDateEdit, QTimeEdit,
    QTableWidget, QTableWidgetItem, QHeaderView, QAbstractItemView, QMessageBox,
    QFileDialog, QDialog, QDialogButtonBox, QFormLayout, QSpinBox, QScrollArea, QSizePolicy, QTabBar)

import prod_native_core as core
import prod_native_api as api

TABS=('Sort','Check','Check Station','Check&Pack','Pack')
PRESETS={'BAU1':'08:00–19:00','BAU2':'19:00–06:00','CP1':'07:00–18:00','CP2':'18:00–06:00'}
VN_TZ=timezone(timedelta(hours=7))


def fmt(value):return f'{value:.2f}' if core.number(value) else '—'

def whole(value):return str(round(value)) if core.number(value) else '—'


def valid_preset(value):
    match=re.fullmatch(r'([0-2]\d:[0-5]\d)[–-]([0-2]\d:[0-5]\d)',str(value).replace(' ',''))
    if not match:return None
    times=match.groups()
    if times[0]==times[1] or any(int(t[:2])>=24 for t in times):return None
    return times


def add_line(layout,label,widget):
    layout.addWidget(QLabel(label));layout.addWidget(widget)


class CheckProdTab(QWidget):
    def __init__(self, bridge, config, save_config, parent=None):
        super().__init__(parent)
        self.bridge,self.config,self.save_config=bridge,config,save_config
        prefs=config.get('prod_settings') or {}
        if not isinstance(prefs,dict):prefs={}
        self.prefs=prefs
        self.tab=prefs.get('tab','Check') if prefs.get('tab') in TABS else 'Check'
        self.mode=prefs.get('shiftMode','BAU') if prefs.get('shiftMode') in ('BAU','CP') else 'BAU'
        self.shift_number=prefs.get('shiftNumber',1) if prefs.get('shiftNumber') in (1,2) else 1
        self.unit=prefs.get('peopleUnit','item_count') if prefs.get('peopleUnit') in ('item_count','order_count') else 'item_count'
        self.presets=dict(PRESETS)
        for key,value in (prefs.get('shiftPresets') or {}).items():
            if key in self.presets and valid_preset(value):self.presets[key]='–'.join(valid_preset(value))
        self.targets={**core.DEFAULT_TARGETS,**{k:v for k,v in (prefs.get('targets') or {}).items() if k in TABS and core.number(v) and v>0}}
        self.order_targets={**core.DEFAULT_ORDER_TARGETS,**{k:v for k,v in (prefs.get('orderTargets') or {}).items() if k in ('Sort','Check','Pack') and core.number(v) and v>0}}
        self.station_names=core.station_table()
        custom=prefs.get('stationsCustom')
        if isinstance(custom,dict) and isinstance(custom.get('CS'),dict) and isinstance(custom.get('PS'),dict):
            self.station_names={k:custom.get(k,self.station_names.get(k,{})) for k in ('CS','PS','SS')}
        self.rows=[];self.cache_key=None;self.busy=False;self.pending=None;self.task=None
        self.sort_ascending=False;self._light_theme=False;self.warnings=[];self._load_key=None
        self.request_log=[];self._retry_count=0;self._last_request=None
        self._highlighted_employee=None
        self._build()
        self.bridge.completed.connect(self._completed)
        self.bridge.sessionChanged.connect(self._state)
        self._state(self.bridge.connected,self.bridge.session_message)
        self._restore_range()
        self._refresh_controls()
        self._render()

    def _save(self):
        self.config['prod_settings']=self.prefs
        if not self.save_config(self.config):self.message.setText('Không lưu được cài đặt Check Prod.')

    def _build(self):
        self.setObjectName('checkProdRoot')
        # IntraCity ThemeController styles child inputs/buttons and calls set_theme()
        # for the Check Prod-specific cards and checked-state palette.
        main=QVBoxLayout(self);main.setContentsMargins(10,8,10,8);main.setSpacing(9)
        # WMS connection status is already displayed in IntraCity's shared header;
        # avoid a second status/title strip above Check Prod controls.
        mode_row=QHBoxLayout();main.addLayout(mode_row)
        self.bau=QPushButton('BAU');self.cp=QPushButton('CP')
        self.item=QPushButton('Item');self.order=QPushButton('Order')
        self.shift1=QPushButton('Ca 1');self.shift2=QPushButton('Ca 2')
        for button in (self.bau,self.cp,self.item,self.order,self.shift1,self.shift2):
            button.setCheckable(True)
            button.setFixedSize(72,34)
            button.setProperty('prodChoice',True)
            button.setCursor(Qt.CursorShape.PointingHandCursor)
        mode_row.addWidget(self.bau);mode_row.addWidget(self.cp);mode_row.addSpacing(18)
        mode_row.addWidget(self.item);mode_row.addWidget(self.order)
        mode_row.addStretch();mode_row.addWidget(self.shift1);mode_row.addWidget(self.shift2)
        self.edit_shifts=QPushButton('⚙ Giờ ca');mode_row.addWidget(self.edit_shifts)
        self.bau.clicked.connect(lambda:self._set_mode('BAU'))
        self.cp.clicked.connect(lambda:self._set_mode('CP'))
        self.item.clicked.connect(lambda:self._set_unit('item_count'))
        self.order.clicked.connect(lambda:self._set_unit('order_count'))
        self.shift1.clicked.connect(lambda:self._set_shift(1))
        self.shift2.clicked.connect(lambda:self._set_shift(2))
        self.edit_shifts.clicked.connect(self._edit_shifts)

        metrics=QFrame();metrics.setObjectName('prodMetrics')
        metrics_l=QGridLayout(metrics);metrics_l.setContentsMargins(5,5,5,5);metrics_l.setSpacing(10)
        self.metric_values={}
        # Native cards; icons and values are painted by Qt and follow IntraCity's theme.
        metric_defs=[('qty','TỔNG SẢN LƯỢNG','📦'),('weighted','TỔNG PROD','▥'),
                     ('avg','PROD TRUNG BÌNH','◷'),('manHour','GIỜ CÔNG','⏱'),
                     ('below','DƯỚI TARGET','⚠')]
        for i,(key,title,glyph) in enumerate(metric_defs):
            cell=QFrame();cell.setObjectName('prodMetricCard');cell.setMinimumHeight(93)
            inner=QHBoxLayout(cell);inner.setContentsMargins(12,12,12,12);inner.setSpacing(13)
            symbol=QLabel(glyph);symbol.setObjectName('prodMetricIcon');symbol.setProperty('tone',key)
            symbol.setAlignment(Qt.AlignmentFlag.AlignCenter);symbol.setFixedSize(50,50)
            symbol.setFont(QFont('Segoe UI Symbol',21))
            inner.addWidget(symbol)
            details=QVBoxLayout();details.setSpacing(6)
            small=QLabel(title);small.setObjectName('prodMetricLabel');details.addWidget(small)
            big=QLabel('—');big.setObjectName('prodMetricValue');details.addWidget(big)
            details.addStretch(1);inner.addLayout(details,1)
            self.metric_values[key]=big
            metrics_l.addWidget(cell,i//3,i%3)
        main.addWidget(metrics)

        # Keep the five work tabs stable; Target actions have their own line so
        # resizing / choosing a tab cannot hide one of the neighbouring tabs.
        work_row=QHBoxLayout();main.addLayout(work_row)
        work_row.addWidget(QLabel('KHU VỰC LÀM VIỆC'))
        self.tab_select=QTabBar();self.tab_select.setObjectName('prodWorkTabs')
        self.tab_select.setExpanding(False)
        self.tab_select.setUsesScrollButtons(False)
        self.tab_select.setSizePolicy(QSizePolicy.Fixed,QSizePolicy.Fixed)
        for name in TABS:self.tab_select.addTab(name)
        self.tab_select.setStyleSheet('QTabBar#prodWorkTabs::tab {min-width:95px; max-width:95px; padding:7px 6px;}')
        work_row.addWidget(self.tab_select)
        work_row.addStretch(1)
        self.tab_select.currentChanged.connect(self._tab_changed)

        target_row=QHBoxLayout();main.addLayout(target_row)
        target_row.addWidget(QLabel('Target Prod'))
        self.target=QLineEdit();self.target.setFixedWidth(74)
        self.target.setToolTip('Target tính theo đơn vị Item hoặc Order đang chọn; không đổi công thức Prod.')
        target_row.addWidget(self.target)
        self.save_target=QPushButton('Lưu Target');self.save_target.clicked.connect(self._save_target)
        target_row.addWidget(self.save_target)
        target_row.addStretch(1)
        self.start=QPushButton('↻ Tải dữ liệu');self.start.clicked.connect(self._start)
        self.start.setObjectName('prodPrimaryAction')
        target_row.addWidget(self.start)

        range_row=QHBoxLayout();main.addLayout(range_row)
        range_row.addWidget(QLabel('KHUNG GIỜ BÁO CÁO'))
        self.date=QDateEdit();self.date.setCalendarPopup(True);self.date.setDisplayFormat('dd/MM/yyyy')
        self.date.setDate(QDate.currentDate());range_row.addWidget(self.date)
        range_row.addWidget(QLabel('Từ'))
        self.from_time=QTimeEdit();self.from_time.setDisplayFormat('HH:mm');range_row.addWidget(self.from_time)
        range_row.addWidget(QLabel('Đến'))
        self.to_time=QComboBox();self.to_time.addItem('Hiện tại','now')
        for h in range(24):
            for m in (0,30):
                label=f'{h:02d}:{m:02d}';self.to_time.addItem(label,label)
        range_row.addWidget(self.to_time)
        range_row.addStretch()
        self.date.dateChanged.connect(self._persist_interval)
        self.from_time.timeChanged.connect(self._persist_interval)
        self.to_time.currentIndexChanged.connect(self._persist_interval)

        filter_row=QHBoxLayout();main.addLayout(filter_row)
        self.show_all=QCheckBox('Hiện tất cả (kể cả đạt Target)');self.show_all.toggled.connect(self._render)
        filter_row.addWidget(self.show_all);filter_row.addStretch()
        self.direction=QPushButton('↓ Cao → Thấp');self.direction.clicked.connect(self._flip)
        filter_row.addWidget(self.direction)
        self.export=QPushButton('Xuất Excel');self.export.clicked.connect(self._export)
        filter_row.addWidget(self.export)
        self.mapping=QPushButton('Station mapping');self.mapping.clicked.connect(self._mapping)
        filter_row.addWidget(self.mapping)

        self.message=QLabel();self.message.setWordWrap(True);self.message.setObjectName('prodMessage')
        main.addWidget(self.message)
        count_row=QHBoxLayout();main.addLayout(count_row)
        self.count=QLabel();count_row.addWidget(self.count);count_row.addStretch()
        self.clear_highlight=QPushButton('Bỏ tô sáng')
        self.clear_highlight.setEnabled(False)
        self.clear_highlight.clicked.connect(self._clear_employee_highlight)
        count_row.addWidget(self.clear_highlight)
        self.table=QTableWidget();self.table.setAlternatingRowColors(True)
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.table.cellClicked.connect(self._employee_clicked)
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
        self.table.horizontalHeader().setStretchLastSection(True)
        main.addWidget(self.table,1)

    def set_theme(self, light):
        # Presentation only: formulas, request payloads and data are untouched.
        self._light_theme=bool(light)
        bg='#f3f6fb' if light else '#1e1e2e'
        card='#ffffff' if light else '#313244'
        border='#dce5f2' if light else '#494b69'
        fg='#182a49' if light else '#e6eafa'
        secondary='#52627c' if light else '#b4c0e5'
        button='#e5ebf5' if light else '#313244'
        active='#2563eb' if light else '#89b4fa'
        active_text='#ffffff' if light else '#11111b'
        hover='#dbe7fb' if light else '#45475a'
        style=f'''
        QWidget#checkProdRoot {{background:{bg}; color:{fg};}}
        QFrame#prodMetrics {{background:transparent; border:none;}}
        QFrame#prodMetricCard {{background:{card}; border:1px solid {border}; border-radius:11px;}}
        QLabel#prodMetricLabel {{color:{secondary}; font-size:11px; font-weight:600;}}
        QLabel#prodMetricValue {{color:{fg}; font-size:22px; font-weight:700;}}
        QLabel#prodMetricIcon {{border-radius:12px; font-size:22px; font-weight:700;}}
        QLabel#prodMetricIcon[tone="qty"] {{background:{'#e0edff' if light else '#254471'};color:{'#1467c0' if light else '#b7d6ff'};}}
        QLabel#prodMetricIcon[tone="weighted"] {{background:{'#def5eb' if light else '#24483d'};color:{'#127a53' if light else '#a6e3a1'};}}
        QLabel#prodMetricIcon[tone="avg"] {{background:{'#eee3ff' if light else '#3b3155'};color:{'#8040c5' if light else '#cba6f7'};}}
        QLabel#prodMetricIcon[tone="manHour"] {{background:{'#fff0df' if light else '#4a3828'};color:{'#be6716' if light else '#fab387'};}}
        QLabel#prodMetricIcon[tone="below"] {{background:{'#ffe9ec' if light else '#492b35'};color:{'#ce3145' if light else '#f4a7b9'};}}
        QPushButton[prodChoice="true"] {{background:{button};color:{fg};border:1px solid {border};border-radius:6px;padding:5px 8px;font-weight:600;}}
        QPushButton[prodChoice="true"]:hover {{background:{hover};}}
        QPushButton[prodChoice="true"]:checked {{background:{active};color:{active_text};border-color:{active};}}
        QPushButton[prodChoice="true"]:disabled {{color:{secondary};}}
        QTabBar#prodWorkTabs::tab {{background:{button};color:{fg};border:1px solid {border};border-radius:5px;padding:7px 6px;min-width:95px;max-width:95px;}}
        QTabBar#prodWorkTabs::tab:selected {{background:{active};color:{active_text};font-weight:700;border-color:{active};}}
        QTabBar#prodWorkTabs::tab:hover:!selected {{background:{hover};}}
        QPushButton#prodPrimaryAction {{background:{active};color:{active_text};border:1px solid {active};font-weight:700;}}
        '''
        self.setStyleSheet(style)
        self._apply_employee_highlight()

    def _state(self,connected,message):
        # Keep connection metadata for diagnostics without duplicating the
        # shared header indicator. No change to any bridge or API behavior.
        self._wms_connected=bool(connected)
        self._wms_connection_message=str(message)

    def _restore_range(self):
        interval=self.prefs.get('intervalPrefs') or {}
        try:
            today=datetime.now(VN_TZ).strftime('%Y-%m-%d')
            day=QDate.fromString(interval.get('date',today),'yyyy-MM-dd')
            if day.isValid():self.date.setDate(day)
            t=QTime.fromString(interval.get('from','07:00'),'HH:mm')
            if t.isValid():self.from_time.setTime(t)
            pos=self.to_time.findData(interval.get('to','now'))
            if pos>=0:self.to_time.setCurrentIndex(pos)
        finally:
            self.tab_select.setCurrentIndex(TABS.index(self.tab))

    def _persist_interval(self,*_):
        if not hasattr(self,'prefs'):return
        self.prefs['intervalPrefs']={'date':self.date.date().toString('yyyy-MM-dd'),
                                    'from':self.from_time.time().toString('HH:mm'),
                                    'to':self.to_time.currentData()}
        self._save()

    def _save_target(self):
        try:value=float(self.target.text())
        except ValueError:value=0
        if not math.isfinite(value) or value<=0:
            QMessageBox.warning(self,'Target','Target phải là số lớn hơn 0.');return
        selected=self.order_targets if self.unit=='order_count' and self.tab in ('Sort','Check','Pack') else self.targets
        selected[self.tab]=value
        self.prefs['targets']=self.targets;self.prefs['orderTargets']=self.order_targets;self._save();self._render()

    def _active_target(self):
        if self.unit=='order_count' and self.tab in ('Sort','Check','Pack'):
            return self.order_targets[self.tab]
        return self.targets[self.tab]

    def _refresh_controls(self):
        for b,active in [(self.bau,self.mode=='BAU'),(self.cp,self.mode=='CP'),
                         (self.item,self.unit=='item_count'),(self.order,self.unit=='order_count'),
                         (self.shift1,self.shift_number==1),(self.shift2,self.shift_number==2)]:
            b.setChecked(active)
        enabled=self.tab in ('Sort','Check','Pack')
        for b in (self.item,self.order):b.setEnabled(enabled)
        self.target.setText(str(self._active_target()))
        self.date.setEnabled(self.tab in ('Sort','Check'))
        self.from_time.setEnabled(self.tab in ('Sort','Check'))
        self.to_time.setEnabled(self.tab in ('Sort','Check'))

    def _set_mode(self,value):
        if self.busy:return
        self.mode=value;self.prefs['shiftMode']=value;self._save();self._apply_shift()

    def _set_shift(self,value):
        if self.busy:return
        self.shift_number=value;self.prefs['shiftNumber']=value;self._save();self._apply_shift()

    def _apply_shift(self):
        times=valid_preset(self.presets[self.mode+str(self.shift_number)])
        if times:
            self.from_time.setTime(QTime.fromString(times[0],'HH:mm'))
            pos=self.to_time.findData(times[1]);self.to_time.setCurrentIndex(pos if pos>=0 else 0)
        self._refresh_controls();self._render()

    def _set_unit(self,value):
        if self.busy or value==self.unit:return
        self.unit=value;self.prefs['peopleUnit']=value;self._save();self.cache_key=None;self._refresh_controls();self._render()

    def _tab_changed(self,index):
        if index<0 or index>=len(TABS):return
        value=TABS[index]
        self.tab=value;self.prefs['tab']=value;self._save();self._refresh_controls();self._render()

    def _edit_shifts(self):
        dialog=QDialog(self);dialog.setWindowTitle('Điều chỉnh khung giờ BAU / CP')
        form=QFormLayout(dialog)
        fields={}
        for key in PRESETS:
            w=QLineEdit(self.presets[key]);fields[key]=w;form.addRow(key,w)
        buttons=QDialogButtonBox(QDialogButtonBox.StandardButton.Save|QDialogButtonBox.StandardButton.Cancel)
        form.addWidget(buttons);buttons.accepted.connect(dialog.accept);buttons.rejected.connect(dialog.reject)
        if not dialog.exec():return
        next_presets={}
        for key,w in fields.items():
            val=valid_preset(w.text())
            if not val:
                QMessageBox.warning(self,'Khung giờ','Cần nhập HH:mm–HH:mm, giờ bắt đầu khác giờ kết thúc.');return
            next_presets[key]='–'.join(val)
        self.presets=next_presets;self.prefs['shiftPresets']=self.presets;self._save();self._apply_shift()

    def _interval(self,now):
        day=self.date.date().toString('yyyy-MM-dd')
        start=self.from_time.time().toString('HH:mm')
        from_dt=datetime.fromisoformat(day+'T'+start+':00').replace(tzinfo=VN_TZ)
        from_ts=int(from_dt.timestamp())
        end=self.to_time.currentData()
        if end=='now':to_ts=now
        else:
            to_ts=int(datetime.fromisoformat(day+'T'+end+':00').replace(tzinfo=VN_TZ).timestamp())
            if to_ts<=from_ts:to_ts+=86400
        if from_ts>now:raise ValueError('Giờ bắt đầu không được ở tương lai.')
        if to_ts>now:to_ts=now
        if to_ts<=from_ts:raise ValueError('Giờ kết thúc phải sau giờ bắt đầu.')
        return {'from':from_ts,'to':to_ts}

    def _load(self, tab, unit, now, email):
        self.warnings=[]
        if tab in ('Sort','Check'):
            interval=self._interval(now)
            items=yield from api.fetch_productivity(interval['from'],interval['to'],
                    43 if tab=='Sort' else 41,unit,'individual',email)
            return core.build_people(items)
        win=core.shift_window(now); hw=core.hour_windows(now)
        if tab=='Pack':
            items=yield from api.fetch_productivity(win['from'],win['to'],42,'item_count','individual',email)
            orders=yield from api.fetch_productivity(win['from'],win['to'],42,'order_count','individual',email)
            try:tasks=yield from api.fetch_pack_tasks()
            except api.WMSRequestError as exc:tasks=[];self.warnings.append(str(exc)+' · thiếu Station')
            prev=yield from self._soft_hours(hw['prev'],42,email)
            cur=yield from self._soft_hours(hw['cur'],42,email,cache=items if hw['cur']['from']==win['from'] and hw['cur']['to']==win['to'] else None)
            return core.merge_rows(items,orders,core.build_pack_station_map(tasks),{'prev':core.hour_map(prev),'cur':core.hour_map(cur)},self.station_names)
        base={'from':win['from'],'to':win['to']}
        stations=yield from api.fetch_productivity(base['from'],base['to'],41,'item_count','station',email)
        if not stations and win['end']!=win['to']:
            alternate=yield from api.fetch_productivity(base['from'],win['end'],41,'item_count','station',email)
            if alternate:stations=alternate;base['to']=win['end']
        if not stations:return []
        try:people=yield from api.fetch_productivity(base['from'],base['to'],41,'item_count','individual',email)
        except api.WMSRequestError as exc:people=[];self.warnings.append('Không có dữ liệu nhân viên: '+str(exc))
        try:tasks=yield from api.fetch_checked_recent(now,core.station_ids(stations))
        except api.WMSRequestError as exc:tasks={'rows':[],'partial':str(exc),'complete':False}
        if tasks.get('partial') and not tasks.get('complete'):self.warnings.append('Thiếu Task gần đây: '+tasks['partial'])
        prev=yield from self._soft_hours(hw['prev'],41,email)
        cur=yield from self._soft_hours(hw['cur'],41,email,cache=people if hw['cur']['from']==base['from'] and hw['cur']['to']==base['to'] else None)
        return core.build_station_rows(stations,people,core.last_operator_by_station(tasks['rows']),
                                      {'prev':core.hour_map(prev),'cur':core.hour_map(cur)},self.station_names)

    def _soft_hours(self,window,task_type,email,cache=None):
        if cache is not None:return cache
        try:return (yield from api.fetch_productivity(window['from'],window['to'],task_type,'item_count','individual',email))
        except api.WMSRequestError as exc:
            self.warnings.append('Thiếu Prod theo giờ: '+str(exc));return []

    def _start(self):
        if self.busy:return
        if not self.bridge.connected:
            QMessageBox.warning(self,'Chưa kết nối','Hãy đăng nhập WMS qua Chrome trước khi tải dữ liệu.');return
        try:
            target=float(self.target.text())
            if not math.isfinite(target) or target<=0:raise ValueError('Target phải lớn hơn 0.')
            if self.tab in ('Sort','Check'):self._interval(int(time.time()))
        except (ValueError,TypeError) as exc:QMessageBox.warning(self,'Dữ liệu không hợp lệ',str(exc));return
        self._save_target()
        self.busy=True;self.start.setEnabled(False);self.start.setText('Đang tải…')
        self.message.setText('Đang tải dữ liệu từ WMS…')
        self._load_key=(self.tab,self.unit)
        self.task=self._load(self.tab,self.unit,int(time.time()),self.prefs.get('loginEmail','').strip())
        self._advance(None,first=True)

    def _advance(self,response,first=False):
        if not self.task:return
        try:
            msg=next(self.task) if first else self.task.send(response)
        except StopIteration as stop:
            self.rows=stop.value or []
            self.cache_key=self._load_key
            self.task=None;self.pending=None;self.busy=False
            self.start.setEnabled(True);self.start.setText('↻ Tải dữ liệu')
            self.message.setText(' '.join(self.warnings) if self.warnings else 'Đã tải dữ liệu.')
            self._render();return
        except Exception as exc:
            self.task=None;self.pending=None;self.busy=False
            self.start.setEnabled(True);self.start.setText('↻ Tải dữ liệu')
            self.message.setText('Không tải được Check Prod: '+str(exc))
            return
        self._last_request=msg;self._retry_count=0
        self._send_request(msg)

    def _send_request(self, msg):
        if not self.task:return
        self.pending='prod-native-'+uuid4().hex
        self.bridge.request(self.pending,json.dumps(msg,ensure_ascii=False))

    def _completed(self,request_id,raw):
        if request_id!=self.pending:return
        self.pending=None
        try:reply=json.loads(raw)
        except (ValueError,TypeError):reply={'success':False,'error':'Phản hồi WMS không hợp lệ.'}
        if not isinstance(reply,dict):reply={'success':False,'error':'Phản hồi WMS không hợp lệ.'}
        error=str(reply.get('error','')) if isinstance(reply,dict) else ''
        if not reply.get('success') and ('HTTP 429' in error or 'HTTP 503' in error) and self._retry_count<4:
            delays=(800,1600,3200,6400)
            delay=delays[self._retry_count];self._retry_count+=1
            self.message.setText(f'WMS giới hạn yêu cầu, thử lại lần {self._retry_count}/4 sau {delay/1000:g}s…')
            QTimer.singleShot(delay,lambda:self._send_request(self._last_request))
            return
        QTimer.singleShot(0,lambda:self._advance(reply))

    def _flip(self):
        self.sort_ascending=not self.sort_ascending
        self.direction.setText('↑ Thấp → Cao' if self.sort_ascending else '↓ Cao → Thấp')
        self._render()

    def _employee_clicked(self,row,col):
        name_col=0 if self.tab in ('Sort','Check','Pack') else 5
        if col!=name_col:return
        item=self.table.item(row,col)
        key=item.data(Qt.ItemDataRole.UserRole) if item else None
        if not key:return
        self._highlighted_employee=None if key==self._highlighted_employee else key
        self.table.clearSelection()
        self._apply_employee_highlight()

    def _clear_employee_highlight(self):
        self._highlighted_employee=None
        self.table.clearSelection()
        self._apply_employee_highlight()

    def _apply_employee_highlight(self):
        # Only decorate table items; keep source rows, totals and export values intact.
        bg=QColor('#fff0ac' if self._light_theme else '#66521c')
        fg=QColor('#302400' if self._light_theme else '#fff3bb')
        for row in range(self.table.rowCount()):
            for col in range(self.table.columnCount()):
                item=self.table.item(row,col)
                if item is None:continue
                active=bool(self._highlighted_employee and
                            item.data(Qt.ItemDataRole.UserRole)==self._highlighted_employee)
                item.setData(Qt.ItemDataRole.BackgroundRole,bg if active else None)
                item.setData(Qt.ItemDataRole.ForegroundRole,fg if active else None)
                font=item.font();font.setBold(active);item.setFont(font)
        self.clear_highlight.setEnabled(self._highlighted_employee is not None)

    def _employee_key(self,row):
        people=self.tab in ('Sort','Check','Pack')
        email=str(row.get('email' if people else 'operatorEmail') or '').strip().casefold()
        if email:return 'email:'+email
        name=str(row.get('name' if people else 'operatorName') or '').strip().casefold()
        return 'name:'+name if name and name not in ('—','-') else None

    def _render(self,*_):
        current=self.cache_key==(self.tab,self.unit)
        if not current:
            self._display_empty()
            return
        target=self._active_target()
        view=core.filtered_view(self.rows,self.tab,target,self.show_all.isChecked(),self.unit,self.sort_ascending)
        metrics=core.summary(self.rows,self.tab,self.unit,len(view['below']))
        for key,value in metrics.items():
            if key in self.metric_values:self.metric_values[key].setText(whole(value) if key in ('qty','below') else fmt(value))
        self.count.setText(f"{len(view['below'])} dưới Target trong tổng {view['total']} · {self.tab}"+
                           (f" · {view['hidden']} Station chưa xác định" if view['hidden'] else ''))
        if self.tab in ('Sort','Check'):
            headers=['Nhân viên','Email', 'Sản lượng '+('Order' if self.unit=='order_count' else 'Item'),'Giờ công','Prod/h','Target']
            make=lambda r:[r.get('name'),r.get('email'),fmt(r.get('items')),fmt(r.get('manHour')),fmt(r.get('productivity')),fmt(target)]
        elif self.tab=='Pack':
            headers=['Nhân viên','Email','Prod (Item)','Prod (Order)','Giờ công','Prod H-1','Prod H','Station ID','Tên Station']
            make=lambda r:[r.get('name'),r.get('email'),fmt(r.get('items')),fmt(r.get('orders')),fmt(r.get('manHour')),
                           fmt(r.get('hPrev')),fmt(r.get('hCur')),r.get('stationId'),r.get('stationName')]
        else:
            headers=['Station ID','Tên Station','Prod Station','Giờ công','Item','Nhân viên','Email','Prod nhân viên','Prod H-1','Prod H']
            make=lambda r:[r.get('stationId'),r.get('stationName'),fmt(r.get('prod')),fmt(r.get('manHour')),
                           whole(r.get('items')),r.get('operatorName'),r.get('operatorEmail'),fmt(r.get('staffProd')),
                           fmt(r.get('hPrev')),fmt(r.get('hCur'))]
        self.table.setUpdatesEnabled(False)
        try:
            self.table.setColumnCount(len(headers));self.table.setHorizontalHeaderLabels(headers)
            self.table.setRowCount(len(view['rows']))
            for row,r in enumerate(view['rows']):
                employee=self._employee_key(r)
                name_col=0 if self.tab in ('Sort','Check','Pack') else 5
                for col,value in enumerate(make(r)):
                    item=QTableWidgetItem(str(value) if value is not None else '—')
                    item.setData(Qt.ItemDataRole.UserRole,employee)
                    if col==name_col and employee:
                        item.setToolTip('Bấm tên để tô sáng tất cả dòng của nhân viên; bấm lại để bỏ tô sáng.')
                    if col and isinstance(value,str) and re.fullmatch(r'\d+\.\d+',value):item.setTextAlignment(Qt.AlignmentFlag.AlignRight|Qt.AlignmentFlag.AlignVCenter)
                    self.table.setItem(row,col,item)
            self._apply_employee_highlight()
        finally:self.table.setUpdatesEnabled(True)

    def _display_empty(self):
        self.table.setRowCount(0)
        self._clear_employee_highlight()
        for value in self.metric_values.values():value.setText('—')
        self.count.setText('Bấm Tải dữ liệu để xem '+self.tab+' · '+('Order' if self.unit=='order_count' else 'Item'))

    def _export(self):
        if self.table.rowCount()==0:
            QMessageBox.information(self,'Xuất Excel','Chưa có danh sách để xuất.');return
        path,_=QFileDialog.getSaveFileName(self,'Xuất Check Prod','Check_Prod_IntraCity_3.4.xlsx','Excel (*.xlsx)')
        if not path:return
        if not path.lower().endswith('.xlsx'):path+='.xlsx'
        from openpyxl import Workbook
        wb=Workbook();ws=wb.active;ws.title='Check Prod'
        ws.append([self.table.horizontalHeaderItem(i).text() for i in range(self.table.columnCount())])
        for i in range(self.table.rowCount()):
            ws.append([self.table.item(i,j).text() if self.table.item(i,j) else '' for j in range(self.table.columnCount())])
        ws.freeze_panes='A2';wb.save(path)

    def _mapping(self):
        choice=QMessageBox(self)
        choice.setWindowTitle('Station mapping')
        choice.setText('Chọn nguồn tên Station:')
        load=choice.addButton('Tải Excel',QMessageBox.ButtonRole.AcceptRole)
        reset=choice.addButton('Dùng bảng mặc định',QMessageBox.ButtonRole.ResetRole)
        choice.addButton('Hủy',QMessageBox.ButtonRole.RejectRole)
        choice.exec()
        if choice.clickedButton() is reset:
            self.station_names=core.station_table()
            self.prefs.pop('stationsCustom',None)
            self._save();self.message.setText('Đã khôi phục Station mapping mặc định.')
            self._render();return
        if choice.clickedButton() is not load:return
        file,_=QFileDialog.getOpenFileName(self,'Chọn Station mapping','','Excel (*.xlsx)')
        if not file:return
        from openpyxl import load_workbook
        try:
            book=load_workbook(file,read_only=True,data_only=True)
            rows=book.worksheets[0].iter_rows(values_only=True)
            table={'CS':{},'PS':{},'SS':{}}
            for row in rows:
                sid=str(row[0]).strip().upper() if len(row)>0 and row[0] is not None else ''
                name=str(row[1]).strip() if len(row)>1 and row[1] is not None else ''
                prefix=sid[:2]
                if re.fullmatch(r'(CS|PS|SS)\d{5}',sid) and name and sid not in table[prefix]:table[prefix][sid]=name
            book.close()
            if not any(table[k] for k in table):raise ValueError('Không có Station hợp lệ trong hai cột A/B.')
        except Exception as exc:
            QMessageBox.warning(self,'Station mapping','Không đọc được file: '+str(exc));return
        self.station_names=table
        self.prefs['stationsCustom']={**table,'fileName':Path(file).name,
                     'cs':len(table['CS']),'ps':len(table['PS']),'ss':len(table['SS'])}
        self._save();self.message.setText('Đã cập nhật Station mapping: '+Path(file).name)
        self._render()
